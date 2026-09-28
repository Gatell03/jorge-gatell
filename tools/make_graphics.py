"""Genera los gráficos propios de la web (todo anónimo, sin imágenes clínicas).

  public/paper.webp      textura de fondo de la web (sin costuras al repetirse)
  public/paper-tile.webp baldosa de papel para el suelo del juego 3D
  public/tape.webp       cinta de carrocero con bordes rasgados
  public/contact.webp    rotulador amarillo "temblando" (botón de contacto)
  public/alba.webp       ilustración de un trazado cefalométrico en la app
  public/generador.webp  miniaturas de la estructura real de plantilla.pptx
  public/f18-thumb.webp, public/drift-thumb.webp  miniaturas de los juegos
  public/og.png          imagen para compartir el enlace (WhatsApp, LinkedIn…)

Uso: python tools/make_graphics.py  (desde la raíz del proyecto)
"""
import math
import random
import re
from pathlib import Path

from PIL import Image, ImageChops, ImageDraw, ImageFilter, ImageFont, ImageOps

ROOT = Path(__file__).resolve().parent.parent
PUBLIC = ROOT / "public"
PLANTILLA = ROOT.parent / "Generador_Estudios_Alba" / "plantilla.pptx"
CARTON = ROOT / "assets-src" / "imagen-carton.jpg"  # foto original del papel arrugado
FONTS = Path("C:/Windows/Fonts")

INK = (31, 27, 22)
YELLOW = (245, 206, 10)


def font(size, bold=False):
    name = "segoeuib.ttf" if bold else "segoeui.ttf"
    try:
        return ImageFont.truetype(str(FONTS / name), size)
    except OSError:
        return ImageFont.load_default()


def bezier(points, steps=40):
    """Curva de Bézier de cualquier grado (De Casteljau)."""
    out = []
    for i in range(steps + 1):
        t = i / steps
        pts = list(points)
        while len(pts) > 1:
            pts = [((1 - t) * a[0] + t * b[0], (1 - t) * a[1] + t * b[1]) for a, b in zip(pts, pts[1:])]
        out.append(pts[0])
    return out


def path(*segments):
    pts = []
    for seg in segments:
        pts += bezier(seg) if len(seg) > 2 else list(seg)
    return pts


# --------------------------------------------------------------------------
# Papel: fondo de la web y baldosa del suelo 3D (a partir de la foto original)
# --------------------------------------------------------------------------
def _paper_blend(img, amount):
    base = Image.new("RGB", img.size, (252, 250, 245))
    return Image.blend(base, ImageChops.multiply(base, img), amount)


def paper_background():
    tex = Image.open(CARTON).convert("RGB")
    tex.thumbnail((2400, 10000), Image.LANCZOS)
    tex = _paper_blend(tex, 0.6)
    # Imagen + su reflejo vertical: al repetirse en vertical no se ve la costura
    tile = Image.new("RGB", (tex.width, tex.height * 2))
    tile.paste(tex, (0, 0))
    tile.paste(ImageOps.flip(tex), (0, tex.height))
    tile.save(PUBLIC / "paper.webp", "WEBP", quality=78, method=6)


def paper_tile():
    src = Image.open(CARTON).convert("RGB").crop((800, 1200, 3200, 3600)).resize((512, 512), Image.LANCZOS)
    t = _paper_blend(src, 0.22)
    tile = Image.new("RGB", (1024, 1024))
    tile.paste(t, (0, 0))
    tile.paste(ImageOps.mirror(t), (512, 0))
    tile.paste(ImageOps.flip(t), (0, 512))
    tile.paste(ImageOps.flip(ImageOps.mirror(t)), (512, 512))
    tile.save(PUBLIC / "paper-tile.webp", "WEBP", quality=80, method=6)


# --------------------------------------------------------------------------
# Cinta adhesiva
# --------------------------------------------------------------------------
def tape():
    rnd = random.Random(7)
    w, h, s = 260, 64, 3  # se dibuja a 3x y se reduce
    W, H = w * s, h * s
    mask = Image.new("L", (W, H), 0)
    d = ImageDraw.Draw(mask)
    # Bordes rasgados: dientes finos e irregulares, como al cortar con los dedos
    left, right = [], []
    y = 0
    while y <= H:
        left.append((rnd.randint(0, 4) * s, y))
        right.append((W - rnd.randint(0, 4) * s, y))
        y += rnd.randint(2, 5) * s
    d.polygon(left + right[::-1], fill=255)

    body = Image.new("RGBA", (W, H))
    px = body.load()
    for yy in range(H):
        shine = 1 + 0.05 * math.cos(yy / H * math.pi)  # brillo muy suave
        for xx in range(W):
            g = rnd.randint(-5, 5)
            px[xx, yy] = (min(255, int(238 * shine) + g), min(255, int(228 * shine) + g),
                          min(255, int(196 * shine) + g), 165)
    body = body.filter(ImageFilter.GaussianBlur(0.6 * s))

    alpha = Image.composite(body.getchannel("A"), Image.new("L", (W, H), 0), mask)
    body.putalpha(alpha)
    body.resize((w, h), Image.LANCZOS).save(PUBLIC / "tape.webp", "WEBP", quality=90, method=6)


# --------------------------------------------------------------------------
# Rotulador amarillo animado (botón de contacto)
# --------------------------------------------------------------------------
def contact():
    w, h, s = 240, 90, 2
    frames = []
    for f in range(8):
        rnd = random.Random(100 + f)
        # Contorno de tinta tembloroso (estilo flipnote, como la estrella)
        outline = []
        for i in range(72):
            a = 2 * math.pi * i / 72
            rx, ry = (w / 2 - 5) * s, (h / 2 - 5) * s
            outline.append((w / 2 * s + math.cos(a) * rx + rnd.uniform(-1.6, 1.6) * s,
                            h / 2 * s + math.sin(a) * ry + rnd.uniform(-1.6, 1.6) * s))
        # Relleno con rotulador: pasadas horizontales que se solapan y tiemblan...
        fill = Image.new("RGBA", (w * s, h * s), (0, 0, 0, 0))
        d = ImageDraw.Draw(fill)
        for row in range(4):
            y0 = (14 + row * 21) * s
            pts = [((x - 10) * s, y0 + rnd.uniform(-3, 3) * s + math.sin(x / 23 + f) * 2 * s)
                   for x in range(0, w + 21, 12)]
            d.line(pts, fill=YELLOW + (240,), width=26 * s, joint="curve")
        # ...recortadas dentro del óvalo (con algo de holgura, como quien se sale un poco)
        mask = Image.new("L", fill.size, 0)
        ImageDraw.Draw(mask).polygon([(x, y) for x, y in outline], fill=255)
        mask = mask.filter(ImageFilter.MaxFilter(5))
        im = fill.copy()
        im.putalpha(ImageChops.multiply(fill.getchannel("A"), mask))
        ImageDraw.Draw(im).line(outline + outline[:2], fill=INK + (220,), width=int(2.4 * s), joint="curve")
        frames.append(im.resize((w, h), Image.LANCZOS))
    frames[0].save(PUBLIC / "contact.webp", "WEBP", save_all=True, append_images=frames[1:],
                   duration=140, loop=0, quality=85, method=6)


# --------------------------------------------------------------------------
# Ilustración de Alba: trazado cefalométrico dentro de la app
# --------------------------------------------------------------------------
def alba():
    s = 2
    W, H = 1600, 1200
    im = Image.new("RGB", (W * s, H * s), (236, 234, 229))
    d = ImageDraw.Draw(im)
    S = lambda p: (p[0] * s, p[1] * s)  # noqa: E731
    Ss = lambda pts: [S(p) for p in pts]  # noqa: E731

    # Ventana de la app
    d.rounded_rectangle(Ss([(60, 60), (1540, 1140)]), radius=22 * s, fill=(250, 250, 248), outline=(210, 207, 200), width=2 * s)
    d.rounded_rectangle(Ss([(60, 60), (1540, 124)]), radius=22 * s, fill=(242, 241, 237))
    d.rectangle(Ss([(60, 100), (1540, 124)]), fill=(242, 241, 237))
    d.line(Ss([(60, 124), (1540, 124)]), fill=(214, 211, 204), width=2 * s)
    for i, c in enumerate([(236, 106, 94), (244, 191, 79), (98, 197, 84)]):
        cx = 100 + i * 30
        d.ellipse(Ss([(cx - 8, 84), (cx + 8, 100)]), fill=c)
    d.text(S((210, 78)), "Proyecto Alba  ·  Trazado cefalométrico", font=font(22 * s, True), fill=(60, 58, 54))

    # Lienzo del trazado (fondo oscuro, como un negatoscopio)
    cx0, cy0 = 90, 150
    d.rounded_rectangle(Ss([(cx0, cy0), (1040, 1110)]), radius=14 * s, fill=(28, 30, 33))
    ox, oy = 90, 190  # desplazamiento del dibujo dentro del lienzo
    P = lambda x, y: (ox + x, oy + y)  # noqa: E731
    bone = (214, 219, 226)
    soft = (140, 190, 235)
    lw = 3 * s

    def curve(pts, col, width=lw):
        d.line(Ss(pts), fill=col, width=width, joint="curve")

    # Contorno craneal y base
    curve(path([P(705, 250), P(730, 60), P(430, -40), P(170, 160)], [P(170, 160), P(120, 300), P(200, 420)]), bone)
    # Silla turca
    curve(path([P(350, 290), P(355, 330), P(405, 335), P(410, 292)]), bone)
    # Órbita
    d.ellipse(Ss([P(570, 250), P(670, 350)]), outline=bone, width=lw)
    # Maxilar y paladar
    curve(path([P(470, 450), P(600, 440), P(700, 445), P(745, 440)]), bone)
    curve(path([P(745, 440), P(700, 460), P(705, 480)], [P(705, 480), P(712, 505), P(722, 525)]), bone)
    # Incisivos
    curve(path([P(690, 450), P(705, 500), P(728, 545)]), bone, 5 * s)
    curve(path([P(672, 622), P(700, 580), P(720, 538)]), bone, 5 * s)
    # Mandíbula
    curve(path([P(715, 560), P(690, 590), P(690, 615)], [P(690, 615), P(705, 650), P(700, 675), P(660, 705)],
               [P(660, 705), P(540, 690), P(420, 640), P(395, 600)], [P(395, 600), P(360, 520), P(345, 430)]), bone)
    # Perfil blando
    curve(path([P(760, 70), P(730, 200), P(745, 265)], [P(745, 265), P(800, 330), P(850, 385), P(775, 440)],
               [P(775, 440), P(800, 470), P(795, 505)], [P(795, 505), P(770, 528), P(790, 560)],
               [P(790, 560), P(750, 600), P(780, 680), P(725, 735)], [P(725, 735), P(640, 760), P(560, 770)]), soft)

    pts = {
        "S": P(380, 312), "N": P(705, 255), "Or": P(655, 345), "Po": P(300, 368),
        "A": P(706, 478), "B": P(692, 608), "Pog": P(702, 660), "Me": P(662, 703),
        "Go": P(398, 600), "ANS": P(745, 440), "PNS": P(470, 450),
    }

    def line_ext(a, b, col, ext=0.25, width=2 * s):
        ax, ay = a
        bx, by = b
        dx, dy = bx - ax, by - ay
        d.line(Ss([(ax - dx * ext, ay - dy * ext), (bx + dx * ext, by + dy * ext)]), fill=col, width=width)

    line_ext(pts["S"], pts["N"], (245, 206, 10), 0.18)          # SN
    line_ext(pts["Po"], pts["Or"], (120, 200, 170), 0.12)       # Frankfort
    line_ext(pts["Go"], pts["Me"], (120, 200, 170), 0.2)        # plano mandibular
    d.line(Ss([pts["N"], pts["A"]]), fill=(245, 206, 10), width=2 * s)
    d.line(Ss([pts["N"], pts["B"]]), fill=(245, 206, 10), width=2 * s)
    d.line(Ss([pts["ANS"], pts["PNS"]]), fill=(170, 170, 180), width=2 * s)

    # Arco del ángulo SNA/SNB en N
    nx, ny = pts["N"]
    d.arc(Ss([(nx - 70, ny - 70), (nx + 70, ny + 70)]), 95, 190, fill=(245, 206, 10), width=2 * s)

    lab = font(19 * s, True)
    for name, (x, y) in pts.items():
        d.ellipse(Ss([(x - 7, y - 7), (x + 7, y + 7)]), fill=(245, 206, 10), outline=(28, 30, 33), width=2 * s)
        tx = x + 12 if name not in ("S", "Go", "Po", "PNS") else x - 12 - len(name) * 12
        d.text(S((tx, y - 30)), name, font=lab, fill=(245, 236, 200))

    # Panel lateral de resultados
    px0 = 1070
    d.text(S((px0, 160)), "Análisis de Steiner", font=font(30 * s, True), fill=INK)
    d.text(S((px0, 204)), "Valores de demostración", font=font(18 * s), fill=(110, 105, 98))
    rows = [("SNA", "82.1°", "82 ± 2"), ("SNB", "79.4°", "80 ± 2"), ("ANB", "2.7°", "2 ± 2"),
            ("SN–GoGn", "31.8°", "32 ± 5"), ("U1–NA", "22.3°", "22 ± 5"), ("L1–NB", "25.6°", "25 ± 5")]
    y = 260
    for name, val, norm in rows:
        d.line(Ss([(px0, y - 14), (1510, y - 14)]), fill=(225, 222, 215), width=s)
        d.ellipse(Ss([(px0, y + 10), (px0 + 12, y + 22)]), fill=(98, 170, 84))
        d.text(S((px0 + 26, y)), name, font=font(22 * s, True), fill=INK)
        d.text(S((px0 + 200, y)), val, font=font(22 * s), fill=INK)
        d.text(S((px0 + 310, y)), norm, font=font(19 * s), fill=(120, 115, 108))
        y += 64
    d.line(Ss([(px0, y - 14), (1510, y - 14)]), fill=(225, 222, 215), width=s)
    # Chips de análisis
    y += 30
    x = px0
    for chip in ["Ricketts", "McNamara", "Jarabak", "Downs", "Tweed"]:
        tw = d.textlength(chip, font=font(18 * s)) / s
        d.rounded_rectangle(Ss([(x, y), (x + tw + 28, y + 40)]), radius=20 * s, outline=(190, 186, 178), width=s)
        d.text(S((x + 14, y + 7)), chip, font=font(18 * s), fill=(70, 66, 60))
        x += tw + 40
        if x > 1440:
            x, y = px0, y + 54
    d.rounded_rectangle(Ss([(px0, 1040), (1510, 1100)]), radius=12 * s, fill=INK)
    d.text(S((px0 + 24, 1055)), "Detectar landmarks con IA", font=font(21 * s, True), fill=(250, 245, 230))

    im.resize((W, H), Image.LANCZOS).save(PUBLIC / "alba.webp", "WEBP", quality=88, method=6)


# --------------------------------------------------------------------------
# Generador: miniaturas desde la estructura real de plantilla.pptx
# --------------------------------------------------------------------------
def generador():
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(PLANTILLA))
    sw, sh = prs.slide_width, prs.slide_height

    def walk(shapes):
        for sp in shapes:
            if sp.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sp.shapes)
            else:
                yield sp

    photo_name = re.compile(r"^(extra|intra|oclusal|modelos|rx|trazado|tabla)_")
    s = 2
    W, H = 1600, 1200
    im = Image.new("RGB", (W * s, H * s), (236, 234, 229))
    d = ImageDraw.Draw(im)
    S = lambda x, y: (int(x * s), int(y * s))  # noqa: E731

    # Ventana
    d.rounded_rectangle([S(60, 60), S(1540, 1140)], radius=22 * s, fill=(250, 250, 248), outline=(210, 207, 200), width=2 * s)
    d.rounded_rectangle([S(60, 60), S(1540, 124)], radius=22 * s, fill=(242, 241, 237))
    d.rectangle([S(60, 100), S(1540, 124)], fill=(242, 241, 237))
    d.line([S(60, 124), S(1540, 124)], fill=(214, 211, 204), width=2 * s)
    for i, c in enumerate([(236, 106, 94), (244, 191, 79), (98, 197, 84)]):
        cx = 100 + i * 30
        d.ellipse([S(cx - 8, 84), S(cx + 8, 100)], fill=c)
    d.text(S(210, 78), f"Presentacion_Nueva.pptx  ·  {len(prs.slides)} diapositivas", font=font(22 * s, True), fill=(60, 58, 54))

    cols, rows = 4, 4
    gx, gy, gap = 110, 170, 30
    tw = (W - 2 * gx - gap * (cols - 1)) / cols
    th = tw * sh / sw
    chosen = [i for i in range(len(prs.slides))][: cols * rows]
    for k, idx in enumerate(chosen):
        slide = prs.slides[idx]
        x0 = gx + (k % cols) * (tw + gap)
        y0 = gy + (k // cols) * (th + gap + 26)
        d.rectangle([S(x0 + 3, y0 + 4), S(x0 + tw + 3, y0 + th + 4)], fill=(215, 212, 205))
        d.rectangle([S(x0, y0), S(x0 + tw, y0 + th)], fill=(255, 255, 255), outline=(200, 197, 190), width=s)
        sx, sy = tw / sw, th / sh
        for sp in walk(slide.shapes):
            if sp.width is None or sp.left is None:
                continue
            l, t = max(0, sp.left) * sx, max(0, sp.top) * sy
            r, b = min(sw, sp.left + sp.width) * sx, min(sh, sp.top + sp.height) * sy
            if r <= l or b <= t:
                continue
            box = [S(x0 + l, y0 + t), S(x0 + r, y0 + b)]
            if sp.shape_type == MSO_SHAPE_TYPE.PICTURE or photo_name.match(sp.name or ""):
                d.rectangle(box, fill=(214, 212, 207), outline=(185, 182, 175), width=1)
                if r - l > 40:
                    d.text(S(x0 + l + 4, y0 + t + 3), sp.name[:16], font=font(9 * s), fill=(95, 92, 86))
            elif sp.has_text_frame and sp.text_frame.text.strip():
                text = sp.text_frame.text.strip()
                generic = text.isupper() and "{{" not in text and len(text) < 28
                if generic and t < th * 0.2:
                    d.text(S(x0 + l + 4, y0 + t + 2), text, font=font(10 * s, True), fill=(40, 40, 40))
                else:  # texto clínico o con nombres: solo líneas grises
                    ly = t + 4
                    while ly < b - 3 and ly < t + 40:
                        d.line([S(x0 + l + 4, y0 + ly), S(x0 + l + (r - l) * 0.85, y0 + ly)], fill=(205, 203, 198), width=3 * s)
                        ly += 7
        d.text(S(x0, y0 + th + 6), f"{idx + 1}", font=font(13 * s), fill=(120, 115, 108))

    im.resize((W, H), Image.LANCZOS).save(PUBLIC / "generador.webp", "WEBP", quality=88, method=6)


# --------------------------------------------------------------------------
# Miniaturas de los juegos
# --------------------------------------------------------------------------
def _paper(w, h, seed=1):
    rnd = random.Random(seed)
    im = Image.new("RGB", (w, h), (243, 238, 226))
    px = im.load()
    for _ in range(w * h // 60):
        x, y = rnd.randrange(w), rnd.randrange(h)
        v = rnd.randint(200, 225)
        px[x, y] = (v, v - 5, v - 16)
    return im


def _cloud(d, x, y, s, width):
    for cx, cy, r in [(-0.9, 0.1, 0.6), (0, -0.2, 0.8), (0.95, 0.15, 0.55)]:
        d.ellipse([x + (cx - r) * s, y + (cy - r) * s, x + (cx + r) * s, y + (cy + r) * s], fill=(251, 248, 240), outline=(74, 66, 56), width=width)
    d.rectangle([x - 1.4 * s, y + 0.1 * s, x + 1.4 * s, y + 0.6 * s], fill=(251, 248, 240))
    d.line([x - 1.5 * s, y + 0.6 * s, x + 1.5 * s, y + 0.6 * s], fill=(74, 66, 56), width=width)


def f18_thumb():
    s = 2
    W, H = 1200, 900
    im = _paper(W * s, H * s, 3)
    d = ImageDraw.Draw(im)
    for x, y, sc in [(260, 170, 70), (900, 120, 55), (700, 330, 45)]:
        _cloud(d, x * s, y * s, sc * s, 3 * s)
    def pencil(wx, end, point, label):
        # Lápiz gigante (mismo diseño que el juego): cuerpo, madera, grafito, virola y goma
        w, tip = 110, 90
        down = point > end
        dr = 1 if down else -1
        body_end = point - tip * dr
        y0, y1 = sorted((end, body_end))
        for c, a, b in [((226, 184, 10), 0, 0.3), ((247, 213, 58), 0.3, 0.7), ((232, 191, 12), 0.7, 1)]:
            d.rectangle([(wx + w * a) * s, y0 * s, (wx + w * b) * s, y1 * s], fill=c)
        d.polygon([(wx * s, body_end * s), ((wx + w) * s, body_end * s), ((wx + w / 2) * s, point * s)], fill=(234, 214, 173), outline=INK, width=3 * s)
        k = 0.38
        d.polygon([((wx + w / 2 - w / 2 * k) * s, (point - tip * k * dr) * s), ((wx + w / 2 + w / 2 * k) * s, (point - tip * k * dr) * s), ((wx + w / 2) * s, point * s)], fill=(46, 42, 38))
        e1 = end + 28 * dr
        f1 = e1 + 30 * dr
        d.rectangle([(wx + 3) * s, min(end, e1) * s, (wx + w - 3) * s, max(end, e1) * s], fill=(231, 167, 154))
        d.rectangle([wx * s, min(e1, f1) * s, (wx + w) * s, max(e1, f1) * s], fill=(207, 201, 189), outline=INK, width=2 * s)
        d.line([wx * s, end * s, wx * s, body_end * s], fill=INK, width=3 * s)
        d.line([(wx + w) * s, end * s, (wx + w) * s, body_end * s], fill=INK, width=3 * s)
        mid = (body_end + f1) / 2
        txt = Image.new("RGBA", (int(240 * s), int(40 * s)), (0, 0, 0, 0))
        ImageDraw.Draw(txt).text((120 * s, 20 * s), f"{label} · HB", font=font(24 * s, True), fill=INK + (200,), anchor="mm")
        txt = txt.rotate(90, expand=True)
        im.paste(txt, (int((wx + w / 2) * s - txt.width / 2), int(mid * s - txt.height / 2)), txt)

    for wx, gap_top, label in [(520, 250, "bug"), (930, 420, "deadline")]:
        pencil(wx, -2, gap_top, label)
        pencil(wx, H + 2, gap_top + 260, "404")
    # Estela
    pts = [(20 + i * 10, 600 - math.sin(i / 5) * 40 - i * 2) for i in range(20)]
    for i in range(1, len(pts)):
        d.line([pts[i - 1][0] * s, pts[i - 1][1] * s, pts[i][0] * s, pts[i][1] * s], fill=(120, 112, 100), width=int((2 + i * 0.3) * s))
    plane = Image.open(PUBLIC / "f18.webp").convert("RGBA").rotate(12, expand=True, resample=Image.BICUBIC)
    plane = plane.resize((int(330 * s), int(330 * s * plane.height / plane.width)), Image.LANCZOS)
    im.paste(plane, (int(200 * s), int(400 * s)), plane)
    # Estrella
    cx, cy, r = 740 * s, 380 * s, 34 * s
    star = [(cx + math.cos(i * math.pi / 4 - math.pi / 2) * (r if i % 2 == 0 else r * 0.3),
             cy + math.sin(i * math.pi / 4 - math.pi / 2) * (r if i % 2 == 0 else r * 0.3)) for i in range(8)]
    d.polygon(star, fill=YELLOW, outline=INK, width=3 * s)
    d.text((745 * s, 40 * s), "12", font=font(90 * s), fill=(180, 172, 160))
    im.resize((W, H), Image.LANCZOS).save(PUBLIC / "f18-thumb.webp", "WEBP", quality=85, method=6)


def drift_thumb():
    s = 2
    W, H = 1200, 900
    im = _paper(W * s, H * s, 9).convert("RGBA")
    marks = Image.new("RGBA", im.size, (0, 0, 0, 0))
    md = ImageDraw.Draw(marks)
    # Donuts y curvas de derrape: dos ruedas = dos trazos paralelos
    for cx, cy, r in [(430, 430, 190), (430, 430, 150), (820, 560, 120), (820, 560, 88)]:
        md.ellipse([(cx - r) * s, (cy - r) * s, (cx + r) * s, (cy + r) * s], outline=INK + (120,), width=14 * s)
    for off in (-22, 22):
        pts = [(80 + t * 10, 800 - t * 6 + math.sin(t / 9) * 60 + off) for t in range(60)]
        md.line([(x * s, y * s) for x, y in pts], fill=INK + (110,), width=13 * s, joint="curve")
    marks = marks.filter(ImageFilter.GaussianBlur(1.2 * s))
    im.alpha_composite(marks)
    d = ImageDraw.Draw(im)
    for x, y in [(430, 430), (360, 380), (500, 380), (360, 480), (500, 480), (820, 560), (980, 300), (1040, 380)]:
        d.rectangle([(x - 18) * s, (y - 18) * s, (x + 18) * s, (y + 18) * s], fill=INK)
        d.ellipse([(x - 14) * s, (y - 14) * s, (x + 14) * s, (y + 14) * s], fill=YELLOW, outline=INK, width=2 * s)
        d.ellipse([(x - 5) * s, (y - 5) * s, (x + 5) * s, (y + 5) * s], fill=(243, 238, 226))
    # Coche visto desde arriba, derrapando
    car = Image.new("RGBA", (200 * s, 380 * s), (0, 0, 0, 0))
    cd = ImageDraw.Draw(car)
    for wx, wy in [(8, 70), (152, 70), (8, 270), (152, 270)]:
        cd.rounded_rectangle([wx * s, wy * s, (wx + 40) * s, (wy + 70) * s], radius=10 * s, fill=(42, 38, 34))
    cd.rounded_rectangle([22 * s, 10 * s, 178 * s, 370 * s], radius=34 * s, fill=INK)
    cd.rounded_rectangle([40 * s, 120 * s, 160 * s, 250 * s], radius=18 * s, fill=(243, 238, 226))
    cd.rounded_rectangle([52 * s, 132 * s, 148 * s, 238 * s], radius=12 * s, fill=INK)
    cd.rectangle([88 * s, 10 * s, 112 * s, 120 * s], fill=YELLOW)
    cd.rectangle([18 * s, 330 * s, 182 * s, 356 * s], fill=YELLOW)
    car = car.rotate(-38, expand=True, resample=Image.BICUBIC)
    shadow = Image.new("RGBA", car.size, (0, 0, 0, 0))
    shadow.putalpha(car.getchannel("A").point(lambda a: a * 0.3))
    shadow = shadow.filter(ImageFilter.GaussianBlur(10 * s))
    im.alpha_composite(shadow, (620 * s, 170 * s))
    im.alpha_composite(car, (600 * s, 150 * s))
    im.convert("RGB").resize((W, H), Image.LANCZOS).save(PUBLIC / "drift-thumb.webp", "WEBP", quality=85, method=6)


# --------------------------------------------------------------------------
# Imagen para compartir el enlace (1200x630, PNG por compatibilidad)
# --------------------------------------------------------------------------
def og_image():
    W, H = 1200, 630
    s = 2
    tile = Image.open(PUBLIC / "paper-tile.webp").convert("RGB").resize((512 * s, 512 * s))
    im = Image.new("RGB", (W * s, H * s))
    for x in range(0, W * s, tile.width):
        for y in range(0, H * s, tile.height):
            im.paste(tile, (x, y))
    d = ImageDraw.Draw(im)
    serif = lambda size: ImageFont.truetype(str(FONTS / "georgia.ttf"), size)  # noqa: E731
    # Estrella animada: primer fotograma
    star = Image.open(PUBLIC / "star.webp")
    star.seek(0)
    star = star.convert("RGBA")
    star.thumbnail((150 * s, 150 * s), Image.LANCZOS)
    im.paste(star, (90 * s, 70 * s), star)
    d.text((100 * s, 250 * s), "INGENIERÍA INFORMÁTICA · ZARAGOZA", font=font(22 * s), fill=(74, 66, 56))
    d.text((92 * s, 290 * s), "Hola, soy Jorge.", font=serif(112 * s), fill=INK)
    # "vivir la experiencia" en negro y amarillo, como en la portada
    label = "vivir la experiencia"
    f = serif(40 * s)
    tw = d.textlength(label, font=f)
    x0, y0 = 100 * s, 460 * s
    box = Image.new("RGBA", (int(tw + 30 * s), int(64 * s)), INK + (255,))
    ImageDraw.Draw(box).text((15 * s, 6 * s), label, font=f, fill=YELLOW)
    box = box.rotate(1.2, expand=True, resample=Image.BICUBIC)
    im.paste(box, (x0, y0), box)
    d.text((W * s - 100 * s, H * s - 60 * s), "jorge-gatell", font=font(22 * s), fill=(74, 66, 56), anchor="rs")
    im.resize((W, H), Image.LANCZOS).save(PUBLIC / "og.png", optimize=True)


if __name__ == "__main__":
    paper_background()
    paper_tile()
    tape()
    contact()
    alba()
    generador()
    f18_thumb()
    drift_thumb()
    og_image()
    for name in ("tape", "contact", "alba", "generador", "f18-thumb", "drift-thumb"):
        p = PUBLIC / f"{name}.webp"
        print(f"{p.stat().st_size / 1024:6.0f} KB  {p.name}")
